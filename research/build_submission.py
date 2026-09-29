"""Submission deck built inside the TechArena 2026 template.

Structure follows the mentor's pitch guidance (problem -> solution -> how it
works -> why better, twice for the two KPIs -> validation) plus the one extra
slide Huawei's e-mail of 2026-09-22 explicitly allows for state of the art
and references, on top of the template's fixed title / QR / thank-you slides.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE

SRC = '../TECHARENA 2026 - Project Template.pptx'
OUT = 'Huawei_TechArena_Submission.pptx'
CHAR = RGBColor(0x3A, 0x3D, 0x40)
MUTED = RGBColor(0x6B, 0x73, 0x78)
BRAND = RGBColor(0xD9, 0x48, 0x0F)

prs = Presentation(SRC)
S = list(prs.slides)                                    # the 7 template slides, original order
LAYOUTS = {l.name: l for l in prs.slide_masters[0].slide_layouts}


def ph(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def drop(shape):
    shape._element.getparent().remove(shape._element)


def set_text(shape, paras, size, bold=False, color=None, align=None, space_after=6, bullets=None):
    tf = shape.text_frame
    tf.word_wrap = True
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color


def add_para(shape, text, size, bold=False, color=None, space_before=14):
    """Appends one more paragraph to a shape's existing text frame (no clearing)."""
    tf = shape.text_frame
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return p


def title(slide, name, text, size):
    t = ph(slide, name)
    t.text_frame.clear()
    r = t.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    return t


def picture(slide, path, l, t, w, h, alt):
    pic = slide.shapes.add_picture(path, l, t, w, h)
    pic._element.nvPicPr.cNvPr.set('descr', alt)
    return pic


def takeaway(slide, text, top=6.34, size=15):
    tb = slide.shapes.add_textbox(Inches(0.92), Inches(top), Inches(11.5), Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = BRAND
    return tb


def footer(slide):
    for s in slide.shapes:
        if s.is_placeholder and s.placeholder_format.type is not None and 'FOOTER' in str(s.placeholder_format.type):
            s.text_frame.paragraphs[0].runs[0].text = 'HUAWEI TECHARENA 2026'
            for r in s.text_frame.paragraphs[0].runs[1:]:
                r.text = ''


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ==================================================================== slide 1  — title
s = S[0]
title(s, 'Title 1', 'AI cuts RF tolerance analysis from 277 days to 15 hours', 30)
sub = s.shapes.add_textbox(Inches(1.2), Inches(5.72), Inches(10.93), Inches(0.5))
sub.text_frame.word_wrap = True
p = sub.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = 'Sub-topic 3  |  A physics-checked AI model that stands in for the full-wave solver'
r.font.size = Pt(17)
r.font.color.rgb = MUTED

# ==================================================================== slide 2  — problem (NEW)
s_problem = prs.slides.add_slide(LAYOUTS['Title and Content'])
title(s_problem, 'Title 1', 'A tolerance study takes months to a century of solver time', 30)
drop(ph(s_problem, 'Content Placeholder 2'))
picture(s_problem, 'figs/deck_problem.png', Inches(0.92), Inches(2.0), Inches(11.5), Inches(3.3),
        'Four figures: one full-wave simulation takes 30 minutes; a sensitivity analysis on the solver needs 13,312 runs, 277 days; '
        'a design search on the solver needs 1.92 million runs, 110 years; only 30 percent of boards pass today, because the 0.3 dB design margin is small against 1.3 dB of factory spread.')
takeaway(s_problem, 'Finding and fixing the causes with the solver alone is not affordable.', top=5.75, size=17)

# ==================================================================== slide 3  — solution (was slide 2)
s = S[1]
title(s, 'Title 1', 'One AI model, trained on 60 simulations, replaces the solver', 30)
body = ph(s, 'Content Placeholder 2')
body.text_frame.clear()
set_text(body, ['Learns how the 11 manufacturing tolerances shape the S-parameters.',
                'Obeys the physics: reciprocity built in, passivity checked on every prediction.',
                '60 simulations are enough: ten random 60-simulation models match the full-data results within about 3 points.'],
         17, space_after=16)
drop(ph(s, 'Picture Placeholder 3'))
picture(s, 'figs/deck_flow.png', Inches(6.75), Inches(2.0), Inches(5.67), Inches(4.76),
        'How it works: 60 full-wave simulations train an AI model (kriging) that predicts every S-parameter in milliseconds, '
        'with reciprocity built in and passivity checked. The model feeds two tasks: sensitivity analysis (KPI 2) and design search (KPI 1).')

# ==================================================================== slide 4  — why better: KPI 2, sensitivity (was slide 3)
s = S[2]
title(s, 'Title 1', 'Thirty simulations find the tolerances that matter, with 5× lower error', 28)
drop(ph(s, 'Content Placeholder 2'))
drop(ph(s, 'Content Placeholder 3'))
picture(s, 'figs/deck_sens.png', Inches(0.92), Inches(2.0), Inches(5.67), Inches(4.25),
        'Bar chart, sensitivity analysis from the same 30 simulations: conventional Monte Carlo error 0.078, this project 0.016, '
        'five times lower. Polynomial chaos and a neural network reach the same accuracy as kriging.')
picture(s, 'figs/deck_drivers.png', Inches(6.75), Inches(2.0), Inches(5.67), Inches(4.25),
        'Leading sensitivity drivers for reflection, coupling and isolation from 45 real runs with 90 percent bootstrap '
        'intervals. Board thickness and copper widths A7 and A3 dominate. All nine intervals contain the value from all 540 runs.')
takeaway(s, 'Board thickness and copper widths A7 and A3 drive most of the variation: the tolerances worth tightening.')

# ==================================================================== slide 5  — why better: KPI 1, design cost (was slide 4)
s = S[3]
title(s, 'Title 1', 'A 110-year design search now takes 30 hours', 30)
for name in ('Text Placeholder 2', 'Text Placeholder 4', 'Content Placeholder 3', 'Content Placeholder 5'):
    drop(ph(s, name))
picture(s, 'figs/deck_design.png', Inches(0.92), Inches(2.0), Inches(11.5), Inches(3.9),
        'Design search over 961 candidate designs, each tested on 2,000 virtual boards. Solver only: 110 years, 1.92 million '
        'simulations one after another. This project: 30 hours, 60 simulations to train the model, then 5 minutes of computing.')
takeaway(s, 'The same 60-simulation model also delivers the sensitivity analysis: no extra simulations needed.', top=6.15, size=16)

# ==================================================================== slide 6  — validation (was slide 5)
s = S[4]
title(s, 'Title 1', 'Four design changes lift the pass rate to 73–85%', 24)
body = ph(s, 'Text Placeholder 3')
body.text_frame.clear()
set_text(body, ['Changes: higher laminate constant; tighter A7, board-thickness and A3 tolerances.',
                "Each prediction was made without that scenario's real boards: 5 of 5 fall within the real range.",
                'Holds when the illustrative limits are halved or tripled (92–97% agreement).'],
         14, space_after=12)
add_para(body, '0 of 734 million predictions break the physics (passivity).',
         15, bold=True, color=BRAND, space_before=10)
drop(ph(s, 'Picture Placeholder 2'))
picture(s, 'figs/deck_stair.png', Inches(5.67), Inches(1.08), Inches(6.75), Inches(5.33),
        'Bar chart of boards passing under stated limits, from the full-curve model retrained without each scenario\'s boards: today 30 percent, '
        'new laminate 47, plus copper width A7 tolerance halved 57, plus board thickness tolerance halved 65, plus copper width A3 '
        'tolerance halved 85. The real boards in each scenario give 28, 43, 54, 63 and 73 percent with intervals that contain every prediction.')

# ==================================================================== slide 7  — state of the art & references (NEW)
s_sota = prs.slides.add_slide(LAYOUTS['Two Content'])
title(s_sota, 'Title 1', 'Established methods, applied with physics checks and real-board validation', 26)
drop(ph(s_sota, 'Content Placeholder 2'))
drop(ph(s_sota, 'Picture Placeholder 3'))
picture(s_sota, 'figs/deck_sota.png', Inches(0.92), Inches(1.75), Inches(11.5), Inches(3.55),
        'Comparison table, common practice versus this project. Surrogate: kriging is used in about a third of RF surrogate studies; '
        'here it is checked against polynomial chaos and a neural network. Output: the literature compresses the response first; here PCA '
        'raised the error 2.8 times, so all 3,012 outputs are predicted. Physics: physics-informed networks need field data; only port data '
        'exists, so reciprocity is built in and passivity checked on 734 million matrices. Sensitivity: closed-form GP Sobol indices are known; '
        'here applied to 11 manufacturing tolerances, leading tolerances found from 30 simulations. Pass rate: GP-assisted Monte Carlo re-checks boards in the '
        'solver; here no extra runs were allowed, so pass rates are checked against held-out real boards.')
citations = [['[1] Oakley & O\'Hagan (2004), J. R. Stat. Soc. B 66(3)',
              '[2] Marrel et al. (2009), Reliab. Eng. Syst. Saf. 94(3)',
              '[3] Le Gratiet, Cannamela & Iooss (2014), SIAM/ASA J. Uncertain. Quantif.',
              '[4] Systematic review, surrogate-based EM design of RF components (2026), Sensors 26(8):2504'],
             ['[5] Kersaudy et al. (2015), J. Comput. Phys. 286:103-117',
              '[6] Khan & Lowther (2022), IEEE Trans. Magn. 58(9)',
              '[7] Fuhrländer & Schöps (2020), J. Math. Ind. 10']]
for col, (x, w) in zip(citations, ((0.92, 6.3), (7.45, 4.97))):
    refs = s_sota.shapes.add_textbox(Inches(x), Inches(5.5), Inches(w), Inches(1.2))
    rtf = refs.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_right = 0
    for i, c in enumerate(col):
        p2 = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        p2.space_after = Pt(2)
        r2 = p2.add_run()
        r2.text = c
        r2.font.size = Pt(10)
        r2.font.color.rgb = CHAR

# ==================================================================== slide 8  — project folder / QR (was slide 6)
s = S[5]
title(s, 'Title 1', 'Project folder', 28)
body = ph(s, 'Text Placeholder 3')
body.text_frame.clear()
set_text(body, ['Scan to open the full material',
                'README: pipeline, evidence and limitations',
                'Code that reproduces every number',
                'Stored results of every study'], 16, space_after=10)
drop(ph(s, 'Picture Placeholder 2'))
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.29), Inches(1.64), Inches(4.5), Inches(4.16))
box.name = 'QR_PLACEHOLDER'
box.fill.background()
box.line.color.rgb = MUTED
box.line.width = Pt(2)
box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
tf = box.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = 'QR code goes here'
r.font.size = Pt(20)
r.font.bold = True
r.font.color.rgb = MUTED
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = 'run insert_qr.py with your folder link'
r2.font.size = Pt(14)
r2.font.color.rgb = MUTED

# ==================================================================== slide 9  — thank you / call to action (was slide 7)
s = S[6]
title(s, 'Title 1', 'THANK YOU', 40)
cta = s.shapes.add_textbox(Inches(1.2), Inches(5.35), Inches(10.93), Inches(1.0))
cta.text_frame.word_wrap = True
p = cta.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = 'From 277 days of solver time to 15 hours – with the physics checked on every prediction.'
r.font.size = Pt(16)
r.font.bold = True
r.font.color.rgb = BRAND
p2 = cta.text_frame.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = '60 simulations  ·  exact sensitivity analysis  ·  pass rates checked on held-out real boards'
r2.font.size = Pt(13)
r2.font.color.rgb = MUTED

# slides added from a layout lack the footer and slide number the template's own slides carry: copy them over
for new in (s_problem, s_sota):
    for shp in S[1].shapes:
        if shp.is_placeholder and shp.name.startswith(('Footer', 'Slide Number')):
            new.shapes._spTree.append(copy.deepcopy(shp._element))

# ==================================================================== reorder: title, problem, solution, kpi2, kpi1, validation, sota, qr, ending
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
# current append order in ids: [title, solution, kpi2, kpi1, validation, qr, ending, problem, sota]
order = [0, 7, 1, 2, 3, 4, 8, 5, 6]
new_ids = [ids[i] for i in order]
for el in ids:
    sldIdLst.remove(el)
for el in new_ids:
    sldIdLst.append(el)

# ==================================================================== presenter notes, in final slide order
NOTES = [
    # 1 title
    'This project addresses Sub-topic 3: applying AI to the modelling of electromagnetic problems. One full-wave simulation '
    'of the supplied coupler takes about 30 minutes, while the two tasks the challenge evaluates, sensitivity analysis and '
    'design, require thousands to millions of simulations. We present a single surrogate model, built from about 60 simulations, '
    'that replaces the solver for both tasks, respects the S-parameter physics on every prediction, and is validated against '
    'boards it has never seen.',
    # 2 problem
    'The figures on this slide use only the solver cost confirmed by Huawei (30 minutes per run) and the supplied dataset. '
    'A standard Saltelli sensitivity analysis requires 13,312 runs, or 277 days of serial solver time; a design search over '
    '961 candidate designs, each evaluated on 2,000 manufacturing realisations, would take about 110 years. The nominal design '
    'meets the illustrative limits, but its 0.25-0.40 dB margin is small compared with the 1.3 dB spread caused by manufacturing '
    'tolerances, so only about 30% of boards pass.',
    # 3 solution
    'The surrogate is a kriging (Gaussian-process) model of the 11 tolerance parameters. Reciprocity is imposed by construction, '
    'and passivity is verified on every predicted S-matrix; Huawei confirmed that enforcing these S-parameter laws qualifies as '
    'physics-informed for this challenge. The validation results use models trained on the 540 supplied designs, excluding the '
    'held-out boards. A separate test with ten random 60-run training sets reproduced the pass-rate results within 1.7 to 3.3 '
    'points and identified the same two leading tolerances in every case. Polynomial-chaos and neural-network surrogates agree '
    'with kriging within 2 points.',
    # 4 KPI 2
    'Left: error of the sensitivity indices when both methods use the same 30 simulations, measured against a reference computed '
    'from all 540 simulations. Conventional Monte Carlo reaches 0.078, the kriging model 0.016, about five times lower; with 45 and '
    '60 simulations the advantage is three to four times. Polynomial chaos and a neural network perform equivalently to kriging, so '
    'the gain is over the conventional approach named in the criterion. Right: the three leading tolerances per metric with 90% '
    'bootstrap intervals from 45 simulations. Board thickness and copper widths A7 and A3 dominate, and every interval contains the '
    'full-data value. These are the tolerances the design changes on slide 6 act on.',
    # 5 KPI 1
    'Serial solver time at 30 minutes per run, conventional workflow versus this pipeline. The design search evaluates 961 '
    'candidate designs, each on 2,000 simulated manufacturing realisations: 1.92 million surrogate evaluations, completed in '
    'about 5 minutes after the 60 training simulations. Even with 100 solver licences running in parallel, the conventional '
    'design search would still take about a year.',
    # 6 validation
    'The four changes act on the tolerances identified on slide 4: a laminate with a higher dielectric constant (4.02), then halved '
    'tolerances on copper width A7, board thickness and copper width A3. For each scenario, the real boards falling inside it were '
    'removed from training, the model was retrained, and its predicted pass rate was compared with those boards. All five predictions '
    'fall within the 95% range of the real boards. The final scenario rests on 30 boards: the model predicts 85%, the real boards show '
    '73%, and the difference is within sampling uncertainty, hence the range 73-85%. The acceptance limits are illustrative, and '
    'pass/fail agreement remains 92-97% when they are halved or tripled. The full-curve model checked all 734 million predicted '
    'S-matrices across every study, and none violated passivity. Remaining limitations: model bias '
    'cannot be separated from sampling noise with the available boards, the stack-up is confidential, and no additional solver '
    'runs were available to confirm the proposed design.',
    # 7 state of the art
    'Kriging is well established for electromagnetic surrogate modelling, and closed-form Sobol indices from Gaussian processes '
    'are known in statistics. The contribution of this work lies in how these methods are applied and verified: physics is '
    'enforced at the port level because no field data were available, the literature\'s default compression (PCA) was tested '
    'and rejected on this data, and every result was validated against real boards held out from training rather than against '
    'additional simulations.',
    # 8 project folder
    'The repository contains the code that reproduces every figure and number in this presentation, the stored results of each '
    'study, and a README describing the pipeline, the validation and the limitations. The Huawei dataset is not included.',
    # 9 closing
    'In summary: 277 days of solver time are reduced to 15 hours, the 110-year design search to about 30 hours, and every one of '
    '734 million predicted S-matrices satisfies the physical constraints.',
]
assert len(NOTES) == len(prs.slides)
for sl, text in zip(prs.slides, NOTES):
    notes(sl, text)

for sl in prs.slides:
    footer(sl)
prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))

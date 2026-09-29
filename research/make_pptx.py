"""Physics teaching deck. One idea per slide, one verified figure per slide."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

INK    = RGBColor(0x12, 0x14, 0x17)
GREY   = RGBColor(0x5C, 0x6B, 0x73)
PALE   = RGBColor(0xC8, 0xD3, 0xD9)
ORANGE = RGBColor(0xE6, 0x9F, 0x00)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
M = 0.62
TITLE_TOP, TITLE_H = 0.34, 0.72
FIG_TOP, FIG_BOT = 1.30, 6.52
FOOT_TOP, FOOT_H = 6.66, 0.56

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def _tb(slide, l, t, w, h, text, size, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def figure_slide(title, image, takeaway):
    s = prs.slides.add_slide(BLANK)
    _tb(s, M, TITLE_TOP, W - 2 * M, TITLE_H, title, 24, True, INK)
    path = "figs/%s.png" % image
    iw, ih = Image.open(path).size
    band_w, band_h = W - 2 * M, FIG_BOT - FIG_TOP
    scale = min(band_w / iw, band_h / ih)
    w, h = iw * scale, ih * scale
    s.shapes.add_picture(path, Inches(M + (band_w - w) / 2),
                         Inches(FIG_TOP + (band_h - h) / 2), Inches(w), Inches(h))
    _tb(s, M, FOOT_TOP, W - 2 * M, FOOT_H, takeaway, 13, False, GREY)
    s.notes_slide.notes_text_frame.text = takeaway


def title_slide(kicker, title, sub):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK
    _tb(s, 1.0, 2.25, W - 2.0, 0.4, kicker, 14, True, ORANGE)
    _tb(s, 1.0, 2.85, W - 2.0, 1.5, title, 40, True, WHITE)
    _tb(s, 1.0, 4.55, W - 2.8, 1.2, sub, 16, False, PALE)


def section(number, text):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK
    _tb(s, 1.0, 2.9, 2.0, 0.5, number, 16, True, ORANGE)
    _tb(s, 1.0, 3.5, W - 2.0, 1.4, text, 32, True, WHITE)


title_slide(
    "HUAWEI TECH ARENA 2026   SUB-TOPIC 3",
    "The physics behind the project",
    "What the device does, what your data measures, and why a fast model of it "
    "answers the two things the competition marks.")

section("PART 1", "What the device actually is")
figure_slide("A three-port coupler", "01_component",
             "Signal enters port 1. Most of it carries on to port 2. A small, known fraction is "
             "tapped off at port 3, so equipment can monitor the signal without interrupting it.")
figure_slide("A radio signal is a wave", "02_wave",
             "Only two properties matter. How big the wave is, and where it sits in its cycle. "
             "Everything your file records is one of those two.")
figure_slide("Why every measurement has two columns", "03_complex",
             "One number cannot carry both amplitude and phase, so the pair is stored as a point "
             "on a plane. The names real and imaginary are historical, not physical.")

section("PART 2", "What your file measures")
figure_slide("Nine measurements, six of them independent", "04_sparams",
             "S followed by two digits. The second digit is the port the signal went into, the "
             "first is the port it came out of. Reciprocity makes three of the nine free.")
figure_slide("Decibels", "05_decibels",
             "A compact way of writing what fraction of the power survived. Always negative for a "
             "component with no power supply, because you never get out more than you put in.")
figure_slide("One board, swept across frequency", "06_sweep",
             "The device behaves differently at every frequency, so each design is measured at 251 "
             "of them. The dip is the resonance the design is tuned around.")
figure_slide("The eleven numbers that describe one board", "16_parameters",
             "Each is a nominal the designer picks plus a tolerance the factory adds. Huawei "
             "supplied the real tolerances, and every study now uses them.")
figure_slide("Why manufacturing tolerance matters", "07_spread",
             "Identical drawings, 540 different boards. Etching, thickness and material all vary, "
             "and some of the resulting boards will fail specification.")

section("PART 3", "The physics we can exploit")
figure_slide("Three laws your measurements already obey", "11_laws",
             "Each one is a constraint the model can respect by construction rather than learn "
             "from scratch. Testing whether each actually helps is part of the contribution.")

section("PART 4", "The project itself")
figure_slide("Why simulation is slow, and the model is not", "08_slow_fast",
             "One simulation takes 30 minutes. The model never solves Maxwell. It learned the "
             "pattern from examples that did, and answers in about a tenth of a millisecond.")
figure_slide("How the whole thing fits together", "09_architecture",
             "Pay for a small number of simulations once, train once, then ask the model millions "
             "of questions that were previously unaffordable.")
figure_slide("Which manufacturing errors actually matter", "10_sensitivity",
             "Under the real tolerances, board thickness and copper width A7 still explain 89 "
             "percent of reflection. For isolation the material constant fell to second place.")
figure_slide("How many boards pass, and how to raise it", "12_yield",
             "A laminate change takes the pass rate from 32 to 45 percent. Halving three "
             "tolerances takes it to 83. Predictions, still to be confirmed in the solver.")
figure_slide("The redesign lever: the laminate", "22_laminate",
             "A higher dielectric constant cuts isolation failures from 24 to 3 percent. The "
             "best value sits at the edge of the data, so higher values need new simulations.")
figure_slide("You do not need all 540 simulations", "14_dataeff",
             "Sixty designs is enough for engineering accuracy. Needing one ninth of the supplied "
             "data is itself an answer to a challenge the brief names.")

section("PART 5", "What Huawei answered")
figure_slide("Physics-informed is not one thing", "13_levels",
             "Your file has no fields, so level three was never reachable. Huawei confirmed "
             "that level two, the physics of the port measurements, counts.")
figure_slide("What Huawei answered, and what it changed", "17_answers",
             "All three answers are now built into the pipeline. The tolerance table changed "
             "two conclusions, and both are reported.")
figure_slide("What the marking scheme rewards", "15_criteria",
             "Serial times at 30 minutes a simulation. Even with 100 running at once, the "
             "sensitivity study alone would take 11 days.")
figure_slide("What to do next, in order", "20_plan",
             "Simulating laminates above 4.02 and confirming the redesign turns a prediction "
             "into a result. The two questions for Huawei lift the last limits on the data.")

prs.save("Huawei_TechArena_Physics.pptx")
print("saved Huawei_TechArena_Physics.pptx")

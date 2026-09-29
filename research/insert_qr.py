"""Put your folder QR code on the "Project folder" slide of the submission deck.

    python insert_qr.py "https://your-shared-folder-link"

Creates Huawei_TechArena_Submission_final.pptx next to the original and leaves the
original untouched. Share the folder as view-only by link, and do not put the Huawei
dataset in it.
"""
import sys
import segno
from pptx import Presentation
from pptx.util import Inches, Emu

SRC = 'Huawei_TechArena_Submission.pptx'
OUT = 'Huawei_TechArena_Submission_final.pptx'

if len(sys.argv) != 2 or not sys.argv[1].startswith(('http://', 'https://')):
    sys.exit('usage: python insert_qr.py "https://your-shared-folder-link"')
url = sys.argv[1]

segno.make(url, error='m').save('qr_folder.png', scale=24, border=2, dark='#3A3D40', light='white')
prs = Presentation(SRC)
slide, box = next((sl, s) for sl in prs.slides for s in sl.shapes if s.name == 'QR_PLACEHOLDER')
side = min(box.width, box.height)
left = box.left + (box.width - side) // 2
top = box.top + (box.height - side) // 2
box._element.getparent().remove(box._element)
pic = slide.shapes.add_picture('qr_folder.png', left, top, side, side)
pic._element.nvPicPr.cNvPr.set('descr', 'QR code linking to the project folder: ' + url)
prs.save(OUT)
print('saved', OUT)
